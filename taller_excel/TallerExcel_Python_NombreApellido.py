"""Lectura de Excel y solución de los cinco retos de la guía evaluativa."""

from pathlib import Path
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from tabulate import tabulate

ARCHIVO = Path(__file__).resolve().parent / "personas.xlsx"
COLUMNAS = ["id", "name", "company", "email", "mac_address"]
ENCABEZADOS = ["#", "Id", "Name", "Company", "Email", "MAC Address"]


def cargar_datos(ruta=ARCHIVO):
    """Selecciona la hoja activa y guarda los registros en una lista de listas."""
    libro = openpyxl.load_workbook(ruta, data_only=True)
    try:
        hoja = libro.active
        encabezados = [hoja.cell(1, columna).value for columna in range(1, 6)]
        if encabezados != COLUMNAS or hoja.max_column != 5:
            raise ValueError("El Excel debe tener: id, name, company, email, mac_address.")

        datos = []
        ids = set()
        # La fila 1 contiene encabezados. El +1 incluye la última persona.
        for numero_fila in range(2, hoja.max_row + 1):
            fila_temporal = []
            for numero_columna in range(1, hoja.max_column + 1):
                fila_temporal.append(hoja.cell(numero_fila, numero_columna).value)
            if all(valor is None for valor in fila_temporal):
                continue
            identificador = str(fila_temporal[0] or "").strip().lower()
            if not identificador:
                raise ValueError(f"La fila {numero_fila} no tiene ID.")
            if identificador in ids:
                raise ValueError(f"ID repetido en la fila {numero_fila}.")
            ids.add(identificador)
            # El consecutivo es adicional al ID original del archivo.
            datos.append([len(datos) + 1] + fila_temporal)
        return datos
    finally:
        libro.close()


def mostrar_tabla(datos, encabezados=ENCABEZADOS):
    if not datos:
        print("No hay registros para mostrar.")
        return
    print(tabulate(datos, headers=encabezados, tablefmt="fancy_grid", missingval="",
                   stralign="left", numalign="center"))


def mostrar_nombre_email(datos):
    # Reto 1: solo Name y Email, sin columnas adicionales.
    mostrar_tabla([[fila[2], fila[4]] for fila in datos], ["Name", "Email"])


def contar_registros(datos):
    # Reto 2: no se cuenta la fila de encabezados.
    print(f"Total de registros: {len(datos)}")


def filtrar_empresa(datos, texto="Group"):
    # Reto 3: no distingue entre mayúsculas y minúsculas.
    return [fila for fila in datos if texto.lower() in str(fila[3] or "").lower()]


def buscar_id(datos, identificador):
    # Reto 4: se busca el ID del Excel, no el consecutivo de la tabla.
    return [fila for fila in datos
            if str(fila[1]).strip().lower() == identificador.strip().lower()]


def menu(datos):
    # Reto 5: mantiene el menú activo hasta seleccionar Salir.
    while True:
        print("\nCONSULTA DE PERSONAS")
        print("1. Ver todos\n2. Ver Name y Email\n3. Contar registros")
        print("4. Filtrar empresas con Group\n5. Buscar por ID\n0. Salir")
        opcion = input("Selecciona una opción: ").strip()
        if opcion == "1":
            mostrar_tabla(datos)
        elif opcion == "2":
            mostrar_nombre_email(datos)
        elif opcion == "3":
            contar_registros(datos)
        elif opcion == "4":
            print("Empresas que contienen Group:")
            mostrar_tabla(filtrar_empresa(datos))
        elif opcion == "5":
            identificador = input("Copia el ID completo de la persona: ").strip()
            if not identificador:
                print("Debes escribir un ID.")
                continue
            encontrados = buscar_id(datos, identificador)
            if encontrados:
                mostrar_tabla(encontrados)
            else:
                print("No se encontró ninguna persona con ese ID.")
        elif opcion == "0":
            print("Consulta finalizada.")
            break
        else:
            print("Opción inválida. Escribe 0, 1, 2, 3, 4 o 5.")


def main():
    try:
        datos = cargar_datos()
    except FileNotFoundError:
        print("No se encontró personas.xlsx. Déjalo junto al archivo Python.")
        return 1
    except (OSError, ValueError, BadZipFile, InvalidFileException) as error:
        print(f"No fue posible leer el Excel: {error}")
        return 1
    if not datos:
        print("El archivo no contiene personas.")
        return 0
    print(f"Archivo cargado correctamente: {len(datos)} personas.")
    try:
        menu(datos)
    except (EOFError, KeyboardInterrupt):
        print("\nConsulta finalizada.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
